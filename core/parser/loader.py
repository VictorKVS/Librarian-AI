#core/parser/loader.py


from typing import BinaryIO, Optional, Union, Dict
import mimetypes
import io
import os
import magic
from pathlib import Path
import logging
from functools import lru_cache
from abc import ABC, abstractmethod
from tempfile import TemporaryDirectory
from subprocess import run, PIPE

try:
    import img2pdf
except ImportError:
    img2pdf = None
try:
    import ocrspace
except ImportError:
    ocrspace = None

# Настройка логирования
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

# Расширение mime-типов
mimetypes.add_type('application/vnd.openxmlformats-officedocument.wordprocessingml.document', '.docx')
mimetypes.add_type('application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', '.xlsx')
mimetypes.add_type('application/vnd.ms-excel.sheet.macroEnabled.12', '.xlsm')
mimetypes.add_type('application/epub+zip', '.epub')

class FileLoader(ABC):
    @abstractmethod
    def load(self, file_input: Union[str, Path, BinaryIO]) -> BinaryIO:
        pass

    @abstractmethod
    def detect_type(self, file_input: Union[str, Path, BinaryIO]) -> str:
        pass

class SmartLoader(FileLoader):
    def __init__(self):
        try:
            self.magic = magic.Magic(mime=True)
        except Exception:
            logger.warning("python-magic not available, using mimetypes fallback")
            self.magic = None

    @lru_cache(maxsize=1024)
    def detect_type(self, file_input: Union[str, Path, BinaryIO]) -> str:
        try:
            if isinstance(file_input, (str, Path)):
                path_type = self._detect_by_path(file_input)
                content_type = self._detect_by_content(file_input)
                return content_type if content_type != 'application/octet-stream' else path_type
            return self._detect_by_content(file_input)
        except Exception as e:
            logger.error(f"Error detecting file type: {e}")
            return 'application/octet-stream'

    def _detect_by_path(self, file_path: Union[str, Path]) -> str:
        mime, _ = mimetypes.guess_type(str(file_path))
        return mime or 'application/octet-stream'

    def _detect_by_content(self, file_input: Union[str, Path, BinaryIO]) -> str:
        if not self.magic:
            return 'application/octet-stream'
        try:
            data = None
            if isinstance(file_input, (str, Path)):
                with open(file_input, 'rb') as f:
                    data = f.read(2048)
            else:
                pos = file_input.tell()
                data = file_input.read(2048)
                file_input.seek(pos)
            return self.magic.from_buffer(data)
        except Exception as e:
            logger.warning(f"Content detection failed: {e}")
            return 'application/octet-stream'

    def load(self, file_input: Union[str, Path, BinaryIO], mode: str = 'rb') -> BinaryIO:
        if isinstance(file_input, (str, Path)):
            if not os.path.exists(file_input):
                raise FileNotFoundError(f"File not found: {file_input}")
            if not os.access(file_input, os.R_OK):
                raise IOError(f"Permission denied: {file_input}")
            stream = open(file_input, mode)
        elif isinstance(file_input, io.IOBase):
            stream = file_input
        else:
            raise ValueError("Unsupported input type")
        return self._wrap_stream(stream, self.detect_type(file_input))

    def _wrap_stream(self, stream: BinaryIO, mime_type: str) -> BinaryIO:
        if mime_type == 'application/gzip':
            import gzip
            return gzip.GzipFile(fileobj=stream)
        if mime_type == 'application/zip':
            from zipfile import ZipFile
            return ZipFile(stream)
        if mime_type.startswith('image') and img2pdf:
            with TemporaryDirectory() as tmp:
                path = os.path.join(tmp, 'img')
                data = stream.read()
                with open(path, 'wb') as f:
                    f.write(data)
                pdf = img2pdf.convert(path)
                return io.BytesIO(pdf)
        return stream

    def extract_pdf_text(self, file_input: Union[str, Path]) -> str:
        try:
            cmd = ["pdftotext", "-layout", str(file_input), "-"]
            output = run(cmd, stdout=PIPE, stderr=PIPE)
            return output.stdout.decode('utf-8', errors='ignore').strip()
        except FileNotFoundError:
            logger.warning("pdftotext not found")
            return ""

    def extract_docx_text(self, file_input: Union[str, Path, BinaryIO]) -> str:
        """Извлечение текста из .docx файла"""
        try:
            from docx import Document
            doc = Document(file_input)
            return "\n".join(p.text for p in doc.paragraphs)
        except ImportError:
            logger.warning("python-docx not installed")
            return ""

    def extract_xlsx_text(self, file_input: Union[str, Path, BinaryIO]) -> str:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(file_input)
            rows = []
            for sheet in wb.worksheets:
                for row in sheet.iter_rows():
                    rows.extend(map(str, [cell.value for cell in row]))
            return "\n".join(rows)
        except ImportError:
            logger.warning("openpyxl not installed")
            return ""

    def perform_ocr(self, file_input: Union[str, Path, BinaryIO]) -> str:
        if ocrspace:
            client = ocrspace.API(api_key='YOUR_API_KEY')
            res = client.ocr_file(file_input)
            return res['ParsedResults'][0]['ParsedText']
        logger.warning("OCRSpace API unavailable")
        return ""

    def get_file_info(self, file_input: Union[str, Path, BinaryIO]) -> Dict:
        info = {'type': self.detect_type(file_input), 'size': 0, 'encoding': None, 'extension': None}
        if isinstance(file_input, (str, Path)):
            info['size'] = os.path.getsize(file_input)
            info['extension'] = Path(file_input).suffix.lower()
        elif isinstance(file_input, io.IOBase):
            pos = file_input.tell()
            file_input.seek(0, os.SEEK_END)
            info['size'] = file_input.tell()
            file_input.seek(pos)
            if info['type'].startswith('text/'):
                try:
                    import chardet
                    sample = file_input.read(1024)
                    file_input.seek(pos)
                    info['encoding'] = chardet.detect(sample)['encoding']
                except ImportError:
                    logger.debug("chardet not available")
        return info

    def _get_file_sample(self, file_input: Union[str, Path, BinaryIO]) -> Optional[bytes]:
        try:
            if isinstance(file_input, (str, Path)):
                with open(file_input, 'rb') as f:
                    return f.read(1024)
            elif isinstance(file_input, io.IOBase):
                pos = file_input.tell()
                data = file_input.read(1024)
                file_input.seek(pos)
                return data
        except:
            return None